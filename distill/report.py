import csv
import logging
import os
from datetime import datetime

log = logging.getLogger(__name__)

# Columns every master-CSV row carries before the T1..Tn breakdown is appended.
MASTER_COLUMNS = [
    "timestamp", "run", "mode", "seed", "classes_per_task", "epochs",
    "dataset", "arch", "loss", "bic_mode",
    "task", "classes_added", "classes_seen", "seen_acc",
    "bic_alpha", "bic_beta", "bic_b_new_mean", "bic_b_old_mean",
    "val_before", "val_before_new", "val_before_old",
    "val_after", "val_after_new", "val_after_old",
    "final", "average_incremental", "forgetting",
]


def _accuracy_rows(results, tasks, name):
    """The triangular table: after task N, how each task seen so far is doing."""
    model = results[name]
    per_task = model["per_task_accuracy"]
    seen = model["seen_accuracy"]
    width = max((len(row) for row in per_task), default=0)

    header = (["after task", "classes added", "classes seen", "all seen (%)"]
              + [f"T{i + 1} (%)" for i in range(width)])
    rows = []
    for t, (acc, row) in enumerate(zip(seen, per_task)):
        classes = tasks[t] if t < len(tasks) else []
        seen_count = sum(len(tasks[i]) for i in range(min(t + 1, len(tasks))))
        # Blanks, not zeros, for tasks that had not arrived yet -- a zero here would
        # read as "got them all wrong" rather than "was not asked".
        padded = list(row) + [None] * (width - len(row))
        rows.append([t + 1, " ".join(str(c) for c in classes), seen_count, acc] + padded)
    return header, rows


def _summary_rows(results, name):
    model = results[name]
    return ["metric", "value"], [
        ["dataset", results.get("dataset")],
        ["architecture", results.get("arch")],
        ["stage-1 loss", results.get("loss")],
        ["bias-correction mode", results.get("bic_mode")],
        ["tasks", len(results.get("tasks", []))],
        ["final accuracy (%)", model["final"]],
        ["average incremental accuracy (%)", model["average_incremental"]],
        ["average forgetting (points)", model["forgetting"]],
    ]


def _bias_rows(results):
    """One row per task that ran stage 2; the first task has no old classes and is absent."""
    header = ["task", "mode", "alpha", "beta", "b: new-class mean", "b: old-class mean",
              "val before (%)", "before new (%)", "before old (%)",
              "val after (%)", "after new (%)", "after old (%)"]
    rows = []
    for rec in results.get("bias_correction", []):
        bias, num_old = rec.get("bias"), rec.get("num_old", 0)
        new_mean = old_mean = None
        if bias:
            new_mean = sum(bias[num_old:]) / max(len(bias) - num_old, 1)
            old_mean = sum(bias[:num_old]) / num_old if num_old else None
        before, after = rec.get("val_before", {}), rec.get("val_after", {})
        rows.append([
            rec.get("task"), rec.get("mode"), rec.get("alpha"), rec.get("beta"),
            new_mean, old_mean,
            before.get("acc"), before.get("new_acc"), before.get("old_acc"),
            after.get("acc"), after.get("new_acc"), after.get("old_acc"),
        ])
    return header, rows


def build_tables(results, name="model"):
    """(sheet name, header, rows) for every table, in the order they should appear."""
    tasks = results.get("tasks", [])
    tables = [("accuracy", *_accuracy_rows(results, tasks, name)),
              ("summary", *_summary_rows(results, name))]
    bias_header, bias_rows = _bias_rows(results)
    if bias_rows:
        tables.append(("bias correction", bias_header, bias_rows))
    return tables


def _write_csv(tables, save_dir, stem):
    paths = []
    for sheet, header, rows in tables:
        path = os.path.join(save_dir, f"{stem}_{sheet.replace(' ', '_')}.csv")
        with open(path, "w", encoding="utf-8", newline="") as fh:
            writer = csv.writer(fh)
            writer.writerow(header)
            writer.writerows(rows)
        paths.append(path)
    return paths


def write_report(results, save_dir, name="model", stem="accuracy"):
    """Write the spreadsheet; no-op without a save dir.  Returns the path written."""
    if not save_dir:
        return None
    if name not in results:
        log.debug("no metrics recorded for %r, skipping the spreadsheet", name)
        return None

    os.makedirs(save_dir, exist_ok=True)
    tables = build_tables(results, name)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        paths = _write_csv(tables, save_dir, stem)
        log.info("openpyxl is not installed; wrote %d CSV files instead of a workbook",
                 len(paths))
        return paths[0] if paths else None

    book = Workbook()
    book.remove(book.active)  # drop the default empty sheet
    for sheet_name, header, rows in tables:
        sheet = book.create_sheet(sheet_name)
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        for i, title in enumerate(header, start=1):
            sheet.column_dimensions[sheet.cell(row=1, column=i).column_letter].width = \
                max(len(str(title)) + 2, 12)

    path = os.path.join(save_dir, f"{stem}.xlsx")
    book.save(path)
    log.info("wrote %s", path)
    return path


def _bias_by_task(results):
    """The stage-2 record for each task, keyed by task number.

    Task 1 has no old classes and never runs stage 2, so it is simply absent; its row
    gets blank calibration columns rather than zeros, which would read as a correction
    that was fitted and came out null.
    """
    by_task = {}
    for rec in results.get("bias_correction", []):
        bias, num_old = rec.get("bias"), rec.get("num_old", 0)
        new_mean = old_mean = None
        if bias:
            new_mean = sum(bias[num_old:]) / max(len(bias) - num_old, 1)
            old_mean = sum(bias[:num_old]) / num_old if num_old else None
        before, after = rec.get("val_before", {}), rec.get("val_after", {})
        by_task[rec.get("task")] = [
            rec.get("alpha"), rec.get("beta"), new_mean, old_mean,
            before.get("acc"), before.get("new_acc"), before.get("old_acc"),
            after.get("acc"), after.get("new_acc"), after.get("old_acc"),
        ]
    return by_task


def master_rows(results, name="model", extra=None, max_tasks=None):
    """One row per increment: the run's settings, that task's accuracy, the calibration
    fitted at it, the run-level summary, and the T1..Tn breakdown.

    `extra` supplies the sweep-level fields the results dict cannot know -- the
    timestamp, the run tag, and the flag values the run was launched with.
    """
    extra = extra or {}
    model = results[name]
    per_task = model["per_task_accuracy"]
    tasks = results.get("tasks", [])
    width = max_tasks or max((len(row) for row in per_task), default=0)
    by_task = _bias_by_task(results)

    header = MASTER_COLUMNS + [f"T{i + 1}" for i in range(width)]
    rows = []
    for t, (acc, breakdown) in enumerate(zip(model["seen_accuracy"], per_task), start=1):
        classes = tasks[t - 1] if t - 1 < len(tasks) else []
        seen_count = sum(len(tasks[i]) for i in range(min(t, len(tasks))))
        # Blank beyond the tasks that have arrived: a zero would read as a task that
        # was evaluated and scored nothing.
        padded = list(breakdown)[:width] + [None] * max(width - len(breakdown), 0)
        rows.append([
            extra.get("timestamp"), extra.get("run"), extra.get("mode"),
            extra.get("seed"), extra.get("classes_per_task"), extra.get("epochs"),
            results.get("dataset"), results.get("arch"), results.get("loss"),
            results.get("bic_mode"),
            t, " ".join(str(c) for c in classes), seen_count, acc,
            *by_task.get(t, [None] * 10),
            model["final"], model["average_incremental"], model["forgetting"],
            *padded,
        ])
    return header, rows


def append_master_csv(results, path, name="model", extra=None, max_tasks=None):
    """Append this run's per-increment rows to a CSV shared across the whole sweep.

    Appending per run rather than at the end means a sweep that dies on run 19 still
    leaves the first 18 on disk.  If the file exists with a different header -- a sweep
    whose --classes-per-task changed the T-column count, say -- the old file is renamed
    aside rather than written into, since a CSV with two different header shapes in it
    is worse than two files.
    """
    if name not in results:
        log.debug("no metrics recorded for %r, nothing to append", name)
        return None

    header, rows = master_rows(results, name, extra, max_tasks)
    if not rows:
        return None

    parent = os.path.dirname(os.path.abspath(path))
    os.makedirs(parent, exist_ok=True)

    write_header = True
    if os.path.exists(path):
        with open(path, encoding="utf-8", newline="") as fh:
            existing = next(csv.reader(fh), None)
        if existing == header:
            write_header = False
        else:
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            rotated = f"{os.path.splitext(path)[0]}.{stamp}.csv"
            os.rename(path, rotated)
            log.warning("%s had a different header; moved it to %s and started a new one",
                        os.path.basename(path), os.path.basename(rotated))

    with open(path, "a", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        if write_header:
            writer.writerow(header)
        writer.writerows(rows)

    log.info("appended %d rows to %s", len(rows), path)
    return path
